"""Generate YuvaIntern Task 6 final project report (DOCX)."""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "YuvaIntern_Task_6_Final_Project_Assessment.docx")

C_DARK = RGBColor(0x1F, 0x4E, 0x79)
C_BLUE = RGBColor(0x2E, 0x74, 0xB5)
C_GREY = RGBColor(0x59, 0x56, 0x59)

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

h1 = doc.styles["Heading 1"]
h1.font.name = "Calibri"
h1.font.size = Pt(16)
h1.font.bold = True
h1.font.color.rgb = C_DARK

h2 = doc.styles["Heading 2"]
h2.font.name = "Calibri"
h2.font.size = Pt(12.5)
h2.font.bold = True
h2.font.color.rgb = C_BLUE

sec = doc.sections[0]
sec.top_margin = Inches(0.9)
sec.bottom_margin = Inches(0.9)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)

FOOT_LABEL = "YuvaIntern Task 6 \u2013 Final Project and Assessment"
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def ftrun(text="", bold=False, italic=False):
    r = fp.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = C_GREY
    r.bold = bold
    r.italic = italic
    return r


r = ftrun()
fld1 = OxmlElement("w:fldChar")
fld1.set(qn("w:fldCharType"), "begin")
r._r.append(fld1)
r = ftrun()
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = "PAGE"
r._r.append(instr)
r = ftrun()
fld2 = OxmlElement("w:fldChar")
fld2.set(qn("w:fldCharType"), "separate")
r._r.append(fld2)
r = ftrun("1")
fld3 = OxmlElement("w:fldChar")
fld3.set(qn("w:fldCharType"), "end")
r._r.append(fld3)
ftrun(f"  |  {FOOT_LABEL}")


def centered(text, size, color=None, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def cover_field(label, value):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(label)
    r1.font.size = Pt(12)
    r1.bold = True
    r2 = p.add_run(value)
    r2.font.size = Pt(12)
    return p


def h(t, lvl=1):
    doc.add_heading(t, level=lvl)


def p(t):
    return doc.add_paragraph(t)


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def dtable(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, hh in enumerate(headers):
        t.rows[0].cells[j].text = hh
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = v
    doc.add_paragraph("")
    return t


def shotbox(title, desc):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(title)
    r1.bold = True
    r1.font.color.rgb = C_DARK
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(desc)
    r2.italic = True
    r2.font.size = Pt(10)
    doc.add_paragraph("")
    return t


# ================= COVER =================
doc.add_paragraph("")
doc.add_paragraph("")
doc.add_paragraph("")
centered("YUVAINTERN INTERNSHIP", 14, C_BLUE, bold=True)
centered("Task 6 \u2013 Final Project and Assessment", 15, C_BLUE, bold=True)
centered("Student Internship & Placement Management System", 26, C_DARK, bold=True)
centered("Capstone Project \u2013 Full-Stack Web Application", 13, C_GREY)
doc.add_paragraph("")
cover_field("Intern: ", "Abhishek Jadhav")
cover_field("Role: ", "Full Stack Developer Intern")
cover_field("Project: ", "Student Internship & Placement Management System")
cover_field("Domain: ", "Web Development / Placement Management")
cover_field("Date: ", "25 October 2026")
doc.add_paragraph("")
centered("Python 3.14 \u00b7 Flask 3.1 \u00b7 SQLite \u00b7 HTML5 \u00b7 CSS3 \u00b7 Vanilla JavaScript \u00b7 pytest", 10, C_GREY, italic=True)
doc.add_page_break()

# ================= TOC =================
h("Table of Contents", 1)
for item in [
    "1. Title Page and Project Information",
    "2. Abstract",
    "3. Introduction",
    "4. Problem Statement and Project Objectives",
    "5. Requirements and Scope",
    "6. Technology Stack and Justification",
    "7. System Architecture and Data Flow",
    "8. Project Planning and Implementation",
    "9. Major Features and Their Operation",
    "10. User Interface and Screenshots",
    "11. Code Quality, Security and Optimization",
    "12. Testing Methodology and Actual Results",
    "13. Installation and User Guide",
    "14. GitHub Repository and Live Demo",
    "15. Limitations and Future Scope",
    "16. Conclusion and Skills Demonstrated",
    "17. Presentation and Demonstration Checklist",
    "Appendix A \u2013 Submission Description (word count)",
]:
    doc.add_paragraph(item)
pn = doc.add_paragraph()
r0 = pn.add_run("Note: ")
r0.font.size = Pt(10)
r1 = pn.add_run("Page numbers are provided in the footer. Update the automatic table of contents in Word if required.")
r1.italic = True
r1.font.size = Pt(10)
doc.add_paragraph("")

# ================= 2. Abstract =================
DESC_LONG = (
    "This Student Internship and Placement Management System is a full-stack web application designed to help students discover internship and entry-level placement opportunities, submit structured applications, track each application's status over time, and steadily improve their employability through skills and resume management, while giving administrators a controlled back-office to publish listings, review applicants, and export operational reports. Students register and log in to a personal dashboard that summarises total applications, applications under review, shortlisted and selected counts, upcoming application deadlines for opportunities they have not yet applied to, and recent status activity drawn from an auditable history table, so progress is visible at a glance without exposing any other student's private data. Opportunity discovery supports keyword search across title, company and description plus filters for location, skill, internship type, remote, on-site and hybrid mode, paid versus unpaid, and deadline ordering, with each listing showing eligibility, required skills, duration, stipend, and application instructions, and with all sample listings explicitly labelled as training data rather than verified live vacancies. Applications are created through an internal form requiring a meaningful cover letter, are blocked at both the application and database level from duplicating the same student-opportunity pair, and move through the Applied, Under Review, Shortlisted, Selected and Rejected states with timestamps and free-text reviewer notes recorded in a separate history table for transparency. Profiles combine contact and education fields with a many-to-many technical-skills inventory and a printable resume view plus an optional strictly validated file upload, and administrators manage users, listings, application reviews, aggregate statistics and filtered CSV exports through server-validated routes protected by role checks. The system is intentionally built on a lightweight Flask and SQLite stack with server-rendered templates, parameterised queries, session-based authentication with per-session CSRF tokens, and an automated pytest suite covering registration, permissions, search, duplicates, status transitions, cross-student access control, CSV export, invalid input, empty states and date validation, making it suitable as a maintainable, secure-by-default internship project that can be installed locally with pip, seeded idempotently, and extended toward notifications, analytics and production-grade deployment."
)

h("2. Abstract", 1)
p("This report documents the design, implementation, testing and documentation of the YuvaIntern Task 6 capstone: a Student Internship & Placement Management System. The project was selected from existing internship work and improved into a complete, presentation-ready capstone. It provides student registration and role-based login, searchable internship listings with an admin publishing workflow, duplicate-safe applications with auditable status transitions, a personal dashboard, skills and printable-resume management with strictly validated uploads, an admin back-office, and filtered CSV reporting. It was built with Flask and SQLite, verified by 59 automated pytest cases (all passing), seeded idempotently, and prepared for GitHub without secrets. Sample listings are synthetic training data. No screenshots, features or test outcomes are fabricated.")

h("3. Introduction", 1)
p("Task 6, titled Final Project and Assessment, is the capstone milestone of the YuvaIntern internship. After completing fundamentals, frontend and backend tasks, it requires a comprehensive project that integrates the skills learned, with proper planning, implementation, testing, optimization, documentation and presentation. This capstone reuses and improves an existing full-stack project rather than starting from scratch, preserving working code and extending it with better documentation, testing evidence and deployment readiness.")
p("Students typically discover internships through scattered posts, spreadsheets and chat forwards, losing track of deadlines and feedback, while placement cells lack a single view of who applied where and with what outcome. This project therefore provides a central portal where students find opportunities, apply uniformly, monitor status and prepare profiles, while admins curate listings and review pipelines with timestamps and exports. Trust matters as much as convenience: students must only see their own data, and uploads must be safe.")

h("4. Problem Statement and Project Objectives", 1)
p("Problem statement: Students and placement cells lack a single, reliable, secure platform to discover internships and placements, submit applications, track status changes, manage profiles and skills, and generate reports. Existing manual processes are error-prone, untraceable and insecure.")
p("Objectives:")
bullets([
    "Provide secure student and admin accounts with session authentication and CSRF protection.",
    "Publish searchable, filterable, deadline-aware listings with an admin publishing workflow.",
    "Accept internal applications, block duplicates, and track five statuses with history.",
    "Show a personal dashboard (counts, deadlines, recent activity) with strict ownership.",
    "Manage skills and education, a printable resume, and an optional safe file upload.",
    "Offer admin user, listing, review and statistics workflows plus filtered CSV reports.",
    "Verify behaviour with automated tests and document only executed results.",
])

h("5. Requirements and Scope", 1)
h("5.1 Functional Requirements", 2)
dtable(["ID", "Requirement", "Status"], [
    ["F1", "Register, login, logout; first user bootstraps admin", "Implemented"],
    ["F2", "Search title/company/skill/location/type; filter mode/paid/deadline", "Implemented"],
    ["F3", "Admin create/edit/publish/archive/delete with date validation", "Implemented"],
    ["F4", "Apply with cover letter; UNIQUE(user, opportunity) blocks duplicates", "Implemented"],
    ["F5", "Statuses Applied / Under Review / Shortlisted / Selected / Rejected + history", "Implemented"],
    ["F6", "Student dashboard counts + deadlines + recent activity (own data only)", "Implemented"],
    ["F7", "Skills CRUD, education profile, printable resume, validated upload", "Implemented"],
    ["F8", "Admin users/listings/reviews/stats; reports filter + CSV export", "Implemented"],
])
h("5.2 Non-Functional Requirements", 2)
dtable(["ID", "Requirement", "Status"], [
    ["N1", "Hashed passwords, parameterised SQL, CSRF, RBAC, IDOR protection", "Implemented"],
    ["N2", "Responsive, accessible forms, empty states, flash messages", "Implemented"],
    ["N3", "pytest coverage of core flows; idempotent seed; indexed queries", "Implemented"],
    ["N4", "No secrets in repository; .env-based configuration", "Implemented"],
])
h("5.3 Scope Boundaries", 2)
p("In scope: local SQLite database, server-rendered UI, authentication, CRUD, search/filter, applications, status history, profiles, skills, resume upload, admin back-office, CSV reports, automated tests, documentation, deployment instructions.")
p("Out of scope (future work): real job-board scraping, email/SMS notifications, payments, production hosting, pagination UI, analytics charts. These are recorded as future enhancements and were not added to keep the project reliable.")

h("6. Technology Stack and Justification", 1)
p("Only tools actually used are listed. No versions are invented.")
dtable(["Tool / Technology", "Purpose", "Version"], [
    ["Python", "Backend runtime", "3.14.4"],
    ["Flask", "Web framework", "3.1.0"],
    ["SQLite", "Relational database", "built-in"],
    ["HTML5", "Page structure", "15 Jinja templates"],
    ["CSS3", "Styling", "static/style.css"],
    ["JavaScript (Vanilla)", "Minor UI behaviour", "static/app.js"],
    ["pytest", "Automated testing", "9.1.1"],
    ["python-docx", "Report generation", "1.2.0"],
    ["Git / GitHub", "Version control", "to be configured"],
])

h("7. System Architecture and Data Flow", 1)
p("Monolithic Flask app with clearly separated route sections (auth, opportunities, applications, profiles, administration, reporting), plus config.py, schema.sql and seed.py. Server-rendered Jinja templates share base.html; one CSS file handles layout; a tiny JS file adds confirmations. Request flow: route \u2192 login_required / admin_required \u2192 server validation \u2192 parameterised SQL in a transaction \u2192 redirect + flash \u2192 render. Sessions store user_id, role and _csrf; cookies are HttpOnly + SameSite=Lax.")
h("7.1 Module Map", 2)
dtable(["File", "Purpose"], [
    ["app.py", "create_app factory, get_db/init_db, CSRF, decorators, all routes, error pages"],
    ["config.py", "SECRET_KEY, DATABASE, UPLOAD_FOLDER, extension whitelist, size cap"],
    ["schema.sql", "DDL for 7 tables with keys, constraints and indexes"],
    ["seed.py", "Idempotent demo data (admin, student, skills, 5 sample listings)"],
    ["templates/ (15)", "Landing, auth, dashboard, discovery, detail, apply, tracking, profile, resume, admin, reports"],
    ["static/", "style.css responsive theme; app.js confirmations"],
    ["tests/", "test_app.py (13) + test_edge.py (46)"],
])
h("7.2 Database Tables", 2)
dtable(["Table", "Purpose"], [
    ["users", "Accounts, roles, hashed passwords"],
    ["student_profiles", "Education, contact, resume text/file"],
    ["skills", "Skill catalogue"],
    ["student_skills", "Many-to-many user-skill with proficiency"],
    ["opportunities", "Listings with status, deadline, mode, paid flag"],
    ["applications", "Student applications with status; UNIQUE(user, opportunity)"],
    ["application_status_history", "Audit trail of status changes"],
])
p("Relationships: User 1\u20141 profile; user N\u2014M skills; user N\u2014M opportunities via applications; application 1\u2014N history. PRAGMA foreign_keys=ON per connection.")

h("8. Project Planning and Implementation", 1)
h("8.1 Milestones", 2)
dtable(["Milestone", "Description", "Status"], [
    ["M1", "Inspect existing project and select capstone", "Completed"],
    ["M2", "Define specification and requirements", "Completed"],
    ["M3", "Improve implementation, validation, security", "Completed"],
    ["M4", "Run and fix automated tests", "Completed (59/59 pass)"],
    ["M5", "Update documentation and README", "Completed"],
    ["M6", "Generate report and submission description", "Completed"],
    ["M7", "GitHub preparation and deployment instructions", "Completed (manual push pending)"],
])
h("8.2 Implementation Notes", 2)
p("The existing project was preserved. Improvements focused on code quality (ruff clean on app.py), honest documentation of the actual 59-test suite, and preparation of the capstone report. No fake functionality or placeholder buttons were added. Inputs are validated on the server, errors are handled with safe pages, and CSRF and ownership checks protect every POST route.")

h("9. Major Features and Their Operation", 1)
bullets([
    "Authentication: registration with validation, login with hashed passwords, logout, first-user-becomes-admin bootstrap.",
    "Opportunity discovery: keyword search and filters (location, type, mode, paid/unpaid, skill) with applied badges.",
    "Admin publishing: create/edit/publish/archive/delete listings with server-side date and enum validation.",
    "Applications: cover-letter form, duplicate blocking at UI and DB, status transitions with history and notes.",
    "Dashboard: totals, under-review, shortlisted, selected, upcoming deadlines, recent activity (own data only).",
    "Profile and skills: education/contact fields, many-to-many skills, printable resume, validated file upload.",
    "Admin back-office: users, listings, reviews, aggregate stats, filtered CSV export.",
    "Reports: status and date filters, CSV download, empty-state handling, role-based scoping.",
])

h("10. User Interface and Screenshots", 1)
p("Pages delivered: landing (hero + counts + featured), register, login, student dashboard (stat cards + deadlines + activity), discovery (filters + cards + applied badges), detail, apply form, my applications, application detail with history, profile (education + skills + upload), printable resume, admin dashboard, users, listing form, review table, reports with CSV. Layout is responsive with consistent navigation, empty states, validation errors and confirmation flashes. Sample listings carry an explicit training-data notice.")
shotbox("[SCREENSHOT 1 \u2013 STUDENT DASHBOARD]", "Dashboard stat cards, upcoming deadlines and recent activity for student@example.com.")
shotbox("[SCREENSHOT 2 \u2013 OPPORTUNITY DISCOVERY]", "Search filters and result cards with applied badges.")
shotbox("[SCREENSHOT 3 \u2013 APPLICATION DETAIL + TESTS]", "Status history; alongside terminal output of pytest 59 passed.")
shotbox("[SCREENSHOT 4 \u2013 PROFILE + PRINTABLE RESUME]", "Skills editor and resume print preview.")
shotbox("[SCREENSHOT 5 \u2013 ADMIN + REPORTS CSV]", "Admin stats, review table and exported applications.csv.")
p("Note: placeholders follow the project report convention. Genuine captures can be pasted before submission; no fabricated screenshots are included.")

h("11. Code Quality, Security and Optimization", 1)
h("11.1 Code Quality", 2)
bullets([
    "Modular route sections inside app.py; shared helpers for validation and forms.",
    "Parameterised SQL throughout; no string-interpolated queries.",
    "Consistent naming, docstrings, and error handling.",
    "ruff check passes on app.py (no lint violations).",
    "59 automated tests covering auth, RBAC, CRUD, search, filters, duplicates, status, IDOR, CSV, validation, empty states, profile, skills and uploads.",
])
h("11.2 Security", 2)
bullets([
    "Werkzeug password hashing; HttpOnly + SameSite=Lax session cookies.",
    "Per-session CSRF token validated on every POST (400 on mismatch).",
    "login_required and admin_required decorators plus per-record ownership checks (non-owner /applications/<id> \u2192 403).",
    "Server-side validation (email regex, lengths, enums, valid_date) and safe error pages.",
    "Resume uploads restricted to .pdf/.doc/.docx/.txt, secure_filename plus random prefix, 2 MB cap, stored outside templates and never executed.",
    "SECRET_KEY from environment; production fail-fast if weak or short.",
])
h("11.3 Optimization", 2)
bullets([
    "Indexes on users.email, opportunities(status/deadline/company/mode), applications(user/opp/status) and history(app).",
    "LIMIT on featured listings, deadlines and activity avoids full scans.",
    "Single SELECT COUNT(*) queries and JOINed list pages avoid N+1 queries.",
    "One transaction per write with UNIQUE constraints as last-line defence.",
    "No premature caching; optimisations applied only where justified.",
])

h("12. Testing Methodology and Actual Results", 1)
p("pytest with a temporary SQLite database per test and CSRF tokens scraped from GET forms (enforcement stays ON). Each functional requirement maps to at least one test.")
dtable(["Suite", "Tests", "Result"], [
    ["tests/test_app.py", "13", "PASS"],
    ["tests/test_edge.py", "46", "PASS"],
    ["Total", "59", "PASS"],
])
p("Command: python -m pytest tests -v. Result: 59 passed in ~74.6s (Windows, Python 3.14.4, pytest 9.1.1).")
p("Representative tests: test_register_and_login, test_invalid_credentials, test_invalid_register, test_admin_can_create_and_edit_opp, test_student_cannot_create_opp, test_search_filter, test_duplicate_application_blocked, test_status_transition_and_history, test_idor_blocked, test_csv_export, test_invalid_opp_date_rejected, test_empty_db_states, test_csrf_enforced, plus 46 edge tests covering auth, opportunity validation, apply edge cases, admin status, reports/CSV, profile, skills, uploads, admin user deletion, search/filter, 404 handling and dashboard.")
p("Seed idempotence verified: running python seed.py twice prints 'exists, skip' for every row.")

h("13. Installation and User Guide", 1)
dtable(["Step", "Command / Action"], [
    ["1", 'cd "student-internship-placement-system"'],
    ["2", "pip install -r requirements.txt"],
    ["3", "python seed.py"],
    ["4", "python app.py  \u2192  http://127.0.0.1:5000"],
    ["5", "python -m pytest tests -v"],
])
p("Demo accounts after seeding: admin@example.com / admin123 (admin), student@example.com / student123 (student). First registered user on a fresh DB becomes admin automatically. Copy .env.example to .env to set a real SECRET_KEY. Troubleshooting: 'no such table' \u2192 run seed.py; CSRF 400 \u2192 submit via rendered forms; port busy \u2192 change port in app.py; uploads land in uploads/ (git-ignored).")

h("14. GitHub Repository and Live Demo", 1)
p("GitHub username: 0xAbhi13. The project is not yet pushed to a public repository from this environment. To publish:")
bullets([
    "git init -b main",
    "git add .",
    "git commit -m \"YuvaIntern Task 6: Final capstone \u2013 Student Internship & Placement Management System\"",
    "git remote add origin https://github.com/0xAbhi13/<repo-name>.git",
    "git push -u origin main",
])
p("Live demo: a Vercel-ready copy exists in student-internship-placement-system-vercel/ (api/index.py, vercel.json, /tmp SQLite). Deploy via Vercel CLI after setting SECRET_KEY. No live URL is claimed until deployment is actually performed.")
p("GitHub Repository URL: [PASTE ACTUAL URL AFTER PUSH]")
p("Live Demo URL: [PASTE ACTUAL URL AFTER DEPLOYMENT]")

h("15. Limitations and Future Scope", 1)
p("Limitations: no email notifications, no full-text ranking, no pagination beyond LIMITs, no avatar handling, screenshots captured manually, single-process dev server, SQLite on /tmp for serverless (ephemeral).")
p("Future scope: pagination and sorting UI, deadline reminders, saved listings, a recruiter role, analytics charts, magic-byte file sniffing, rate limiting, production WSGI with migrations, CI running pytest on every push, and accessibility audit improvements.")

h("16. Conclusion and Skills Demonstrated", 1)
p("The capstone successfully delivers a complete full-stack placement management system: secure auth, searchable listings, duplicate-safe applications with history, personal and admin dashboards, skills and resume workflows, filtered CSV reports, 59/59 passing tests, an idempotent seed and honest documentation. The project is ready for local demo, GitHub publication and Vercel deployment.")
p("Skills demonstrated: requirements analysis, project planning, Flask routing and blueprints-style modularisation, SQLite schema design, authentication and authorisation, CSRF and IDOR protection, server-side validation, Jinja templating, responsive CSS, file upload handling, automated testing with pytest, debugging, documentation, report generation, presentation planning.")

h("17. Presentation and Demonstration Checklist", 1)
bullets([
    "Explain the problem and target users (students and placement cells).",
    "Show architecture and data flow (Flask \u2192 SQLite \u2192 Jinja templates).",
    "Demonstrate registration/login and role-based access.",
    "Demonstrate opportunity search and filters.",
    "Demonstrate application submission and duplicate blocking.",
    "Demonstrate status transition and history.",
    "Demonstrate dashboard and reports/CSV export.",
    "Demonstrate profile, skills and resume upload.",
    "Show pytest results (59 passed) and seed idempotence.",
    "Discuss security (CSRF, IDOR, validation, uploads).",
    "Discuss limitations and future scope.",
    "Provide GitHub and (if deployed) live demo links.",
])

h("Appendix A \u2013 Submission Description Word Count", 1)
p(DESC_LONG)
pn3 = doc.add_paragraph()
rW = pn3.add_run(f"Word count of description: {len(DESC_LONG.split())} words.")
rW.italic = True

doc.core_properties.title = "YuvaIntern Task 6 \u2013 Final Project and Assessment"
doc.core_properties.author = "Abhishek Jadhav"
doc.core_properties.subject = "Student Internship & Placement Management System"

doc.save(OUT)
body_words = len(" ".join([x.text for x in doc.paragraphs]).split())
print(f"Saved {OUT} | body words={body_words} | desc words={len(DESC_LONG.split())}")
